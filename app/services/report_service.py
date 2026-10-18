"""Audit exceptions report generation service (openpyxl).

Produces multi-sheet Excel (.xlsx) workpapers with:
- Summary: Executive audit scorecard, run traceability, and configured settings.
- Flagged Transactions: Complete line-item exceptions with scores and plain-English reasons.
- Rule Counts: Breakdown across audit rules R1 to R6.
- Benford: Leading-digit distribution and conformity metrics.
- Review Log: Chronological audit disposition and review history.

Includes CSV/Excel formula injection neutralization per docs/SECURITY.md.
"""

from __future__ import annotations

import io
import logging
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import desc, select
from sqlalchemy.orm import Session, joinedload

from app.core.exceptions import ConflictError
from app.db.models.flagged_transaction import FlaggedTransaction
from app.db.models.review import Review
from app.db.models.transaction import Transaction
from app.db.models.user import User
from app.repositories.upload_repo import get_upload_by_id
from app.services.analysis_service import get_run_detail

logger = logging.getLogger(__name__)

# Design System Colors
NAVY_PRIMARY = "1B2A33"
SLATE_ACCENT = "4A5B65"
HEADER_FILL_COLOR = "1B2A33"
SECTION_FILL_COLOR = "E9EFF4"
ZEBRA_FILL_COLOR = "F8FAFC"
BORDER_COLOR = "D9D9D9"

# Rule metadata dictionary for readable Rule Counts sheet
RULE_METADATA = {
    "R1": {
        "name": "Self-approval",
        "category": "Segregation of Duties",
        "points": 40,
        "description": "Transaction created and approved by the same user ID (case-insensitive).",
    },
    "R2": {
        "name": "Duplicate payment",
        "category": "Cash Disbursement",
        "points": 40,
        "description": "Identical vendor and amount posted within configured duplicate window.",
    },
    "R3": {
        "name": "Missing approver",
        "category": "Authorization Control",
        "points": 30,
        "description": "Transaction lacks an designated approver (null, blank, or placeholder).",
    },
    "R4": {
        "name": "Just below approval limit",
        "category": "Policy Evasion",
        "points": 40,
        "description": "Amount falls within threshold window just under single-signoff limit.",
    },
    "R5": {
        "name": "Round amount",
        "category": "Supporting Indicator",
        "points": 10,
        "description": "Amount is an exact multiple of configured round threshold values.",
    },
    "R6": {
        "name": "Weekend / after-hours",
        "category": "Supporting Indicator",
        "points": 10,
        "description": (
            "Transaction posted on Saturday, Sunday, or outside standard business hours."
        ),
    },
}


def sanitize_formula_injection(value: Any) -> Any:
    """Neutralize formula injection by prefixing text starting with =, +, -, @ with '.

    Per docs/SECURITY.md Section 8, any string cell whose leading character
    could trigger Excel formula execution is escaped.

    Args:
        value: Any cell value.

    Returns:
        Sanitized value safe for spreadsheet export.
    """
    if value is None:
        return ""
    if isinstance(value, (bool, int, float, Decimal, datetime, date)):
        return value
    val_str = str(value)
    # Check if stripping leading whitespace begins with formula trigger character
    stripped = val_str.lstrip()
    if stripped and stripped[0] in ("=", "+", "-", "@"):
        return f"'{val_str}"
    return val_str


def _create_styles() -> dict[str, Any]:
    """Build shared reusable openpyxl styles."""
    thin_border = Border(
        left=Side(style="thin", color=BORDER_COLOR),
        right=Side(style="thin", color=BORDER_COLOR),
        top=Side(style="thin", color=BORDER_COLOR),
        bottom=Side(style="thin", color=BORDER_COLOR),
    )
    return {
        "title_font": Font(name="Calibri", size=15, bold=True, color=NAVY_PRIMARY),
        "subtitle_font": Font(name="Calibri", size=10, italic=True, color="5F6C74"),
        "section_font": Font(name="Calibri", size=11, bold=True, color=NAVY_PRIMARY),
        "section_fill": PatternFill(
            start_color=SECTION_FILL_COLOR, end_color=SECTION_FILL_COLOR, fill_type="solid"
        ),
        "header_font": Font(name="Calibri", size=10, bold=True, color="FFFFFF"),
        "header_fill": PatternFill(
            start_color=HEADER_FILL_COLOR, end_color=HEADER_FILL_COLOR, fill_type="solid"
        ),
        "cell_font": Font(name="Calibri", size=10),
        "cell_bold": Font(name="Calibri", size=10, bold=True),
        "thin_border": thin_border,
        "align_center": Alignment(horizontal="center", vertical="center"),
        "align_left": Alignment(horizontal="left", vertical="center"),
        "align_right": Alignment(horizontal="right", vertical="center"),
        "align_wrap": Alignment(horizontal="left", vertical="center", wrap_text=True),
    }


def _auto_fit_columns(ws: Any, min_width: int = 12, max_width: int = 55) -> None:
    """Adjust worksheet column widths according to content length."""
    ws.views.sheetView[0].showGridLines = True
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = cell.value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, max_width), min_width)


def _build_summary_sheet(
    ws: Any,
    run: Any,
    upload: Any,
    flagged_items: list[FlaggedTransaction],
    styles: dict[str, Any],
) -> None:
    """Populate Summary worksheet with executive scorecard, metadata, and settings."""
    ws.title = "Summary"
    ws.views.sheetView[0].showGridLines = True

    # 1. Title Banner
    ws["A1"] = "AuditLens — Exceptions & Anomaly Detection Workpaper Report"
    ws["A1"].font = styles["title_font"]
    ws.row_dimensions[1].height = 24

    ws["A2"] = (
        f"Generated: {datetime.now(UTC).strftime('%Y-%m-%d %H:%M:%S UTC')} | "
        "Strict Auditor Review Working Document"
    )
    ws["A2"].font = styles["subtitle_font"]
    ws.row_dimensions[2].height = 18

    # 2. Traceability & Run Overview
    ws["A4"] = "Run & Dataset Traceability"
    ws["A4"].font = styles["section_font"]
    ws["A4"].fill = styles["section_fill"]
    ws["B4"] = ""
    ws["B4"].fill = styles["section_fill"]
    ws.row_dimensions[4].height = 20

    summary_dict = run.summary or {}
    duration = summary_dict.get("duration_seconds", 0.0)
    signer_info = "Unsigned (Pending Manager Sign-off)"
    if run.signed_off_at:
        signer_name = run.signer.full_name or run.signer.email if run.signer else "Manager"
        signer_info = (
            f"Signed off by {signer_name} on {run.signed_off_at.strftime('%Y-%m-%d %H:%M:%S UTC')}"
        )

    overview_rows = [
        ("Analysis Run ID", run.id),
        ("Dataset Filename", sanitize_formula_injection(upload.original_filename)),
        ("Dataset Upload Date", upload.created_at.strftime("%Y-%m-%d %H:%M:%S UTC")),
        ("Run Execution Status", run.status.upper()),
        (
            "Run Execution Date",
            run.started_at.strftime("%Y-%m-%d %H:%M:%S UTC") if run.started_at else "N/A",
        ),
        ("Pipeline Duration", f"{duration:.2f} seconds"),
        ("Manager Sign-Off Status", sanitize_formula_injection(signer_info)),
    ]

    curr_row = 5
    for label, val in overview_rows:
        ws.cell(row=curr_row, column=1, value=label).font = styles["cell_bold"]
        ws.cell(row=curr_row, column=1).border = styles["thin_border"]
        ws.cell(row=curr_row, column=2, value=val).font = styles["cell_font"]
        ws.cell(row=curr_row, column=2).border = styles["thin_border"]
        curr_row += 1

    # 3. Audit Exception Metrics
    curr_row += 1
    ws.cell(row=curr_row, column=1, value="Audit Exceptions Summary").font = styles["section_font"]
    ws.cell(row=curr_row, column=1).fill = styles["section_fill"]
    ws.cell(row=curr_row, column=2, value="").fill = styles["section_fill"]
    ws.row_dimensions[curr_row].height = 20
    curr_row += 1

    total_accepted = upload.row_count_accepted or 0
    total_flagged = len(flagged_items)
    flag_pct = (total_flagged / total_accepted * 100.0) if total_accepted > 0 else 0.0

    scoring_sum = summary_dict.get("scoring", {})
    by_level = scoring_sum.get("by_level", {})
    benford_sum = summary_dict.get("benford", {})
    benford_label = (
        benford_sum.get("overall", {}).get("conformity_label", "Not evaluated")
        if benford_sum.get("overall")
        else "Insufficient data"
    )

    metric_rows = [
        ("Total Ingested Transactions", upload.row_count_total or 0),
        ("Accepted Valid Transactions", total_accepted),
        ("Total Flagged Exceptions", total_flagged),
        ("Flagged Rate (%)", f"{flag_pct:.2f}%"),
        ("High Risk Items (Score >= 70)", by_level.get("High", by_level.get("high", 0))),
        ("Medium Risk Items (Score 50-69)", by_level.get("Medium", by_level.get("medium", 0))),
        ("Low Risk Items (Score 30-49)", by_level.get("Low", by_level.get("low", 0))),
        ("Benford's Law Conformity", sanitize_formula_injection(benford_label)),
    ]

    for label, val in metric_rows:
        ws.cell(row=curr_row, column=1, value=label).font = styles["cell_bold"]
        ws.cell(row=curr_row, column=1).border = styles["thin_border"]
        ws.cell(row=curr_row, column=2, value=val).font = styles["cell_font"]
        ws.cell(row=curr_row, column=2).border = styles["thin_border"]
        curr_row += 1

    # 4. Settings & Thresholds Configuration
    curr_row += 1
    ws.cell(
        row=curr_row, column=1, value="Detection Rules & Thresholds Configuration"
    ).font = styles["section_font"]
    ws.cell(row=curr_row, column=1).fill = styles["section_fill"]
    ws.cell(row=curr_row, column=2, value="").fill = styles["section_fill"]
    ws.row_dimensions[curr_row].height = 20
    curr_row += 1

    settings = run.settings or {}
    setting_items = [
        ("Single Approval Limit (₹)", f"{settings.get('approval_limit', 100000.0):,.2f}"),
        (
            "Limit Evasion Window",
            f"{float(settings.get('approval_limit_window', 0.10)) * 100:.0f}%",
        ),
        ("Duplicate Window (Days)", f"{settings.get('duplicate_window_days', 7)} days"),
        (
            "Round Amount Modulos",
            ", ".join(str(m) for m in settings.get("round_modulos", [1000, 10000, 100000])),
        ),
        (
            "Standard Business Hours",
            (
                f"{settings.get('business_hours_start', '08:00')} - "
                f"{settings.get('business_hours_end', '20:00')}"
            ),
        ),
        ("ML Contamination Factor", f"{settings.get('ml_contamination', 0.02)}"),
        ("Flag Score Threshold", f"{settings.get('flag_threshold', 30)}"),
        ("Medium Risk Threshold", f"{settings.get('score_medium_threshold', 50)}"),
        ("High Risk Threshold", f"{settings.get('score_high_threshold', 70)}"),
        ("Model Random Seed", f"{settings.get('random_seed', 42)}"),
    ]

    for label, val in setting_items:
        ws.cell(row=curr_row, column=1, value=label).font = styles["cell_bold"]
        ws.cell(row=curr_row, column=1).border = styles["thin_border"]
        ws.cell(row=curr_row, column=2, value=sanitize_formula_injection(val)).font = styles[
            "cell_font"
        ]
        ws.cell(row=curr_row, column=2).border = styles["thin_border"]
        curr_row += 1

    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 46


def _build_flagged_sheet(
    ws: Any,
    flagged_items: list[FlaggedTransaction],
    styles: dict[str, Any],
) -> None:
    """Populate Flagged Transactions worksheet with line item exceptions."""
    ws.title = "Flagged Transactions"
    ws.freeze_panes = "A2"

    headers = [
        "Transaction ID",
        "Posting Date",
        "Amount (₹)",
        "Vendor ID",
        "Department",
        "Account Code",
        "Created By",
        "Approved By",
        "Description",
        "Risk Score",
        "Risk Level",
        "Rule Points",
        "ML Points",
        "Triggered Rules",
        "Reasons",
        "Duplicate Group ID",
        "Review Status",
    ]

    ws.row_dimensions[1].height = 26
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.alignment = styles["align_center"]
        cell.border = styles["thin_border"]

    if not flagged_items:
        # Zero-flag edge case
        ws.row_dimensions[2].height = 22
        cell = ws.cell(
            row=2, column=1, value="No transactions were flagged for review in this analysis run."
        )
        cell.font = styles["cell_bold"]
        cell.alignment = styles["align_left"]
        _auto_fit_columns(ws)
        return

    for row_idx, item in enumerate(flagged_items, start=2):
        ws.row_dimensions[row_idx].height = 20
        txn = item.transaction

        # Transaction attributes
        posting_date_str = (
            txn.posting_date.strftime("%Y-%m-%d %H:%M:%S") if txn.posting_date else ""
        )
        triggered_str = ", ".join(item.triggered_rules or [])
        reasons_str = "; ".join(str(r) for r in (item.reasons or []))
        dup_group = str(item.duplicate_group_id or "")

        row_values = [
            (sanitize_formula_injection(txn.transaction_id), styles["align_center"], None),
            (sanitize_formula_injection(posting_date_str), styles["align_center"], None),
            (float(txn.amount), styles["align_right"], "#,##0.00"),
            (sanitize_formula_injection(txn.vendor_id), styles["align_center"], None),
            (sanitize_formula_injection(txn.department or ""), styles["align_left"], None),
            (sanitize_formula_injection(txn.account_code or ""), styles["align_center"], None),
            (sanitize_formula_injection(txn.created_by or ""), styles["align_center"], None),
            (sanitize_formula_injection(txn.approved_by or ""), styles["align_center"], None),
            (sanitize_formula_injection(txn.description or ""), styles["align_left"], None),
            (int(item.risk_score), styles["align_right"], "0"),
            (sanitize_formula_injection(item.risk_level.title()), styles["align_center"], None),
            (int(item.rule_points), styles["align_right"], "0"),
            (int(item.ml_points), styles["align_right"], "0"),
            (sanitize_formula_injection(triggered_str), styles["align_center"], None),
            (sanitize_formula_injection(reasons_str), styles["align_wrap"], None),
            (sanitize_formula_injection(dup_group), styles["align_center"], None),
            (
                sanitize_formula_injection(item.review_status.replace("_", " ").title()),
                styles["align_center"],
                None,
            ),
        ]

        for col_idx, (val, alignment, num_format) in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = styles["cell_font"]
            cell.alignment = alignment
            cell.border = styles["thin_border"]
            if num_format:
                cell.number_format = num_format

    _auto_fit_columns(ws)


def _build_rule_counts_sheet(
    ws: Any,
    run: Any,
    flagged_items: list[FlaggedTransaction],
    styles: dict[str, Any],
) -> None:
    """Populate Rule Counts worksheet with R1-R6 triggered breakdown."""
    ws.title = "Rule Counts"
    ws.freeze_panes = "A2"

    headers = [
        "Rule Code",
        "Rule Name",
        "Audit Category",
        "Rule Points",
        "Triggered Count",
        "Description",
    ]

    ws.row_dimensions[1].height = 26
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.alignment = styles["align_center"]
        cell.border = styles["thin_border"]

    summary_dict = run.summary or {}
    counts_from_summary = summary_dict.get("rules", {}).get("counts_per_rule", {})

    # Fallback to counting from flagged items if summary lacks counts
    rule_counts: dict[str, int] = {}
    for r_code in RULE_METADATA:
        if r_code in counts_from_summary:
            rule_counts[r_code] = int(counts_from_summary[r_code])
        else:
            rule_counts[r_code] = sum(
                1 for item in flagged_items if r_code in (item.triggered_rules or [])
            )

    curr_row = 2
    for r_code, meta in sorted(RULE_METADATA.items()):
        ws.row_dimensions[curr_row].height = 22
        ws.cell(row=curr_row, column=1, value=r_code).alignment = styles["align_center"]
        ws.cell(row=curr_row, column=2, value=meta["name"]).alignment = styles["align_left"]
        ws.cell(row=curr_row, column=3, value=meta["category"]).alignment = styles["align_left"]
        ws.cell(row=curr_row, column=4, value=meta["points"]).alignment = styles["align_right"]
        ws.cell(row=curr_row, column=5, value=rule_counts.get(r_code, 0)).alignment = styles[
            "align_right"
        ]
        ws.cell(row=curr_row, column=6, value=meta["description"]).alignment = styles["align_left"]

        for c_idx in range(1, 7):
            cell = ws.cell(row=curr_row, column=c_idx)
            cell.font = styles["cell_font"]
            cell.border = styles["thin_border"]

        curr_row += 1

    _auto_fit_columns(ws)


def _build_benford_sheet(
    ws: Any,
    run: Any,
    styles: dict[str, Any],
) -> None:
    """Populate Benford worksheet with digit distribution and vendor metrics."""
    ws.title = "Benford"
    ws.views.sheetView[0].showGridLines = True

    ws["A1"] = "Benford's Law First-Digit Analysis"
    ws["A1"].font = styles["title_font"]
    ws.row_dimensions[1].height = 24

    summary_dict = run.summary or {}
    benford = summary_dict.get("benford", {})
    overall = benford.get("overall")

    if not overall or not overall.get("has_sufficient_data"):
        ws["A3"] = (
            "Benford's Law analysis not available: insufficient transaction sample size "
            "(requires at least 500 positive transaction amounts)."
        )
        ws["A3"].font = styles["cell_bold"]
        _auto_fit_columns(ws)
        return

    # Section 1: Overall Conformity Assessment
    ws["A3"] = "Overall Conformity Assessment"
    ws["A3"].font = styles["section_font"]
    ws["A3"].fill = styles["section_fill"]
    ws["B3"] = ""
    ws["B3"].fill = styles["section_fill"]
    ws.row_dimensions[3].height = 20

    conformity_stats = [
        ("Analysis Sample Size", overall.get("sample_size", 0)),
        ("Conformity Level", sanitize_formula_injection(overall.get("conformity_label", "N/A"))),
        ("Mean Absolute Deviation (MAD)", f"{overall.get('mad', 0.0):.4f}"),
        ("Chi-Square Statistic", f"{overall.get('chi_square', 0.0):.2f}"),
        ("P-Value", f"{overall.get('p_value', 0.0):.4f}"),
    ]

    curr_row = 4
    for label, val in conformity_stats:
        ws.cell(row=curr_row, column=1, value=label).font = styles["cell_bold"]
        ws.cell(row=curr_row, column=1).border = styles["thin_border"]
        ws.cell(row=curr_row, column=2, value=val).font = styles["cell_font"]
        ws.cell(row=curr_row, column=2).border = styles["thin_border"]
        curr_row += 1

    # Section 2: First-Digit Frequency Distribution Table
    curr_row += 2
    ws.cell(row=curr_row, column=1, value="Digit Frequencies (1 - 9)").font = styles["section_font"]
    ws.cell(row=curr_row, column=1).fill = styles["section_fill"]
    for c_i in range(2, 6):
        ws.cell(row=curr_row, column=c_i, value="").fill = styles["section_fill"]
    ws.row_dimensions[curr_row].height = 20
    curr_row += 1

    table_headers = [
        "Leading Digit",
        "Observed Count",
        "Observed %",
        "Expected Benford %",
        "Difference %",
    ]
    ws.row_dimensions[curr_row].height = 24
    for col_idx, h in enumerate(table_headers, start=1):
        cell = ws.cell(row=curr_row, column=col_idx, value=h)
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.alignment = styles["align_center"]
        cell.border = styles["thin_border"]

    obs_counts = overall.get("observed_counts", {})
    obs_freqs = overall.get("observed_frequencies", {})
    exp_freqs = overall.get("expected_frequencies", {})

    curr_row += 1
    for d in range(1, 10):
        d_str = str(d)
        cnt = int(obs_counts.get(d_str, obs_counts.get(d, 0)))
        obs_p = float(obs_freqs.get(d_str, obs_freqs.get(d, 0.0)))
        exp_p = float(exp_freqs.get(d_str, exp_freqs.get(d, 0.0)))
        diff_p = obs_p - exp_p

        ws.row_dimensions[curr_row].height = 20
        ws.cell(row=curr_row, column=1, value=d).alignment = styles["align_center"]
        ws.cell(row=curr_row, column=2, value=cnt).alignment = styles["align_right"]
        ws.cell(row=curr_row, column=3, value=obs_p).alignment = styles["align_right"]
        ws.cell(row=curr_row, column=3).number_format = "0.0%"
        ws.cell(row=curr_row, column=4, value=exp_p).alignment = styles["align_right"]
        ws.cell(row=curr_row, column=4).number_format = "0.0%"
        ws.cell(row=curr_row, column=5, value=diff_p).alignment = styles["align_right"]
        ws.cell(row=curr_row, column=5).number_format = "+0.0%;-0.0%;0.0%"

        for c_i in range(1, 6):
            cell = ws.cell(row=curr_row, column=c_i)
            cell.font = styles["cell_font"]
            cell.border = styles["thin_border"]
        curr_row += 1

    # Section 3: Vendor Conformity Evaluation (if any)
    vendors = benford.get("vendors", {})
    if vendors:
        curr_row += 2
        ws.cell(row=curr_row, column=1, value="Per-Vendor Benford Conformity").font = styles[
            "section_font"
        ]
        ws.cell(row=curr_row, column=1).fill = styles["section_fill"]
        for c_i in range(2, 6):
            ws.cell(row=curr_row, column=c_i, value="").fill = styles["section_fill"]
        ws.row_dimensions[curr_row].height = 20
        curr_row += 1

        v_headers = ["Vendor ID", "Sample Size", "MAD", "Chi-Square", "Conformity Status"]
        ws.row_dimensions[curr_row].height = 24
        for col_idx, h in enumerate(v_headers, start=1):
            cell = ws.cell(row=curr_row, column=col_idx, value=h)
            cell.font = styles["header_font"]
            cell.fill = styles["header_fill"]
            cell.alignment = styles["align_center"]
            cell.border = styles["thin_border"]

        curr_row += 1
        # Sort vendors by MAD descending
        sorted_vendors = sorted(
            vendors.items(),
            key=lambda item: float(item[1].get("mad", 0.0)),
            reverse=True,
        )
        for v_id, v_data in sorted_vendors:
            ws.row_dimensions[curr_row].height = 20
            ws.cell(
                row=curr_row, column=1, value=sanitize_formula_injection(str(v_id))
            ).alignment = styles["align_center"]
            ws.cell(
                row=curr_row, column=2, value=int(v_data.get("sample_size", 0))
            ).alignment = styles["align_right"]
            ws.cell(row=curr_row, column=3, value=float(v_data.get("mad", 0.0))).alignment = styles[
                "align_right"
            ]
            ws.cell(row=curr_row, column=3).number_format = "0.0000"
            ws.cell(
                row=curr_row, column=4, value=float(v_data.get("chi_square", 0.0))
            ).alignment = styles["align_right"]
            ws.cell(row=curr_row, column=4).number_format = "0.00"
            ws.cell(
                row=curr_row,
                column=5,
                value=sanitize_formula_injection(str(v_data.get("conformity_label", "N/A"))),
            ).alignment = styles["align_center"]

            for c_i in range(1, 6):
                cell = ws.cell(row=curr_row, column=c_i)
                cell.font = styles["cell_font"]
                cell.border = styles["thin_border"]
            curr_row += 1

    _auto_fit_columns(ws)


def _build_review_sheet(
    ws: Any,
    reviews: list[Review],
    styles: dict[str, Any],
) -> None:
    """Populate Review Log worksheet with full audit disposition trail."""
    ws.title = "Review Log"
    ws.freeze_panes = "A2"

    headers = [
        "Logged At (UTC)",
        "Transaction ID",
        "Reviewer",
        "Review Decision",
        "Risk Level",
        "Auditor Comment / Justification",
    ]

    ws.row_dimensions[1].height = 26
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = styles["header_font"]
        cell.fill = styles["header_fill"]
        cell.alignment = styles["align_center"]
        cell.border = styles["thin_border"]

    if not reviews:
        ws.row_dimensions[2].height = 22
        cell = ws.cell(
            row=2,
            column=1,
            value="No audit review notes or status changes recorded yet for this run.",
        )
        cell.font = styles["cell_bold"]
        cell.alignment = styles["align_left"]
        _auto_fit_columns(ws)
        return

    for row_idx, rev in enumerate(reviews, start=2):
        ws.row_dimensions[row_idx].height = 20
        flagged = rev.flagged_transaction
        txn = flagged.transaction if flagged else None
        txn_id = txn.transaction_id if txn else f"Flagged #{rev.flagged_id}"
        risk_lvl = flagged.risk_level.title() if flagged else "N/A"
        reviewer_name = (
            rev.reviewer.full_name or rev.reviewer.email
            if rev.reviewer
            else f"User {rev.reviewer_id}"
        )
        created_str = rev.created_at.strftime("%Y-%m-%d %H:%M:%S") if rev.created_at else ""

        status_formatted = rev.status.replace("_", " ").title()

        row_vals = [
            (sanitize_formula_injection(created_str), styles["align_center"]),
            (sanitize_formula_injection(txn_id), styles["align_center"]),
            (sanitize_formula_injection(reviewer_name), styles["align_left"]),
            (sanitize_formula_injection(status_formatted), styles["align_center"]),
            (sanitize_formula_injection(risk_lvl), styles["align_center"]),
            (sanitize_formula_injection(rev.comment or ""), styles["align_wrap"]),
        ]

        for col_idx, (val, alignment) in enumerate(row_vals, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = styles["cell_font"]
            cell.alignment = alignment
            cell.border = styles["thin_border"]

    _auto_fit_columns(ws)


def generate_excel_report(
    db: Session,
    run_id: int,
    current_user: User,
) -> tuple[bytes, str]:
    """Generate multi-sheet Excel audit workpaper report for an analysis run.

    Enforces ownership permissions (404 for other users' runs unless manager)
    and validates completed run status (409 if not completed).

    Args:
        db: Active database session.
        run_id: Analysis run primary key ID.
        current_user: Authenticated user.

    Returns:
        tuple[bytes, str]: (Raw xlsx binary bytes, recommended filename).

    Raises:
        NotFoundError: If run not found or unowned.
        ConflictError: If run status is not 'completed'.
    """
    run = get_run_detail(db, run_id, current_user)

    if run.status != "completed":
        raise ConflictError(
            code="run_not_completed",
            message=(
                f"Cannot generate report for run with status '{run.status}'. Run must be completed."
            ),
        )

    upload = get_upload_by_id(db, run.upload_id)
    if not upload:
        raise ConflictError(
            code="upload_missing",
            message=f"Associated upload for run {run_id} is missing.",
        )

    # Query all flagged transactions for this run with joined transactions
    flagged_stmt = (
        select(FlaggedTransaction)
        .join(Transaction, FlaggedTransaction.transaction_pk == Transaction.id)
        .where(FlaggedTransaction.run_id == run_id)
        .options(joinedload(FlaggedTransaction.transaction))
        .order_by(desc(FlaggedTransaction.risk_score), desc(Transaction.posting_date))
    )
    flagged_items = list(db.scalars(flagged_stmt).unique().all())

    # Query all reviews for this run
    review_stmt = (
        select(Review)
        .join(FlaggedTransaction, Review.flagged_id == FlaggedTransaction.id)
        .join(Transaction, FlaggedTransaction.transaction_pk == Transaction.id)
        .join(User, Review.reviewer_id == User.id)
        .where(FlaggedTransaction.run_id == run_id)
        .options(
            joinedload(Review.flagged_transaction).joinedload(FlaggedTransaction.transaction),
            joinedload(Review.reviewer),
        )
        .order_by(desc(Review.created_at))
    )
    reviews = list(db.scalars(review_stmt).unique().all())

    wb = openpyxl.Workbook()
    styles = _create_styles()

    # 1. Summary sheet (reusing the default active sheet)
    ws_summary = wb.active
    _build_summary_sheet(ws_summary, run, upload, flagged_items, styles)

    # 2. Flagged Transactions sheet
    ws_flagged = wb.create_sheet(title="Flagged Transactions")
    _build_flagged_sheet(ws_flagged, flagged_items, styles)

    # 3. Rule Counts sheet
    ws_rules = wb.create_sheet(title="Rule Counts")
    _build_rule_counts_sheet(ws_rules, run, flagged_items, styles)

    # 4. Benford sheet
    ws_benford = wb.create_sheet(title="Benford")
    _build_benford_sheet(ws_benford, run, styles)

    # 5. Review Log sheet
    ws_reviews = wb.create_sheet(title="Review Log")
    _build_review_sheet(ws_reviews, reviews, styles)

    output = io.BytesIO()
    wb.save(output)
    xlsx_bytes = output.getvalue()

    filename = f"auditlens_report_run_{run_id}.xlsx"
    logger.info(
        "Generated Excel report for run %d (%d bytes, %d flagged items)",
        run_id,
        len(xlsx_bytes),
        len(flagged_items),
    )
    return xlsx_bytes, filename
