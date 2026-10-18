"""Integration tests for Excel audit workpaper report export (Phase 6)."""

import io
from datetime import UTC, datetime

import openpyxl
import pandas as pd
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.db.models.analysis_run import AnalysisRun
from app.db.models.flagged_transaction import FlaggedTransaction
from app.db.models.transaction import Transaction
from app.db.models.upload import Upload
from app.db.models.user import User
from app.repositories.review_repo import create_review
from app.services.analysis_service import execute_analysis_pipeline


def _create_sample_ledger_csv() -> io.BytesIO:
    """Create a sample ledger with known anomalies for report generation."""
    data = [
        {
            "transaction_id": "TXN-NORM-01",
            "posting_date": "2026-01-12 11:00:00",
            "amount": 15000.0,
            "vendor_id": "VEND-10",
            "department": "Finance",
            "account_code": "ACC-101",
            "created_by": "alice",
            "approved_by": "bob",
            "description": "Standard software subscription",
        },
        {
            "transaction_id": "TXN-R1-01",
            "posting_date": "2026-01-14 10:15:00",
            "amount": 85000.0,
            "vendor_id": "VEND-20",
            "department": "Finance",
            "account_code": "ACC-101",
            "created_by": "eve",
            "approved_by": "eve",
            "description": "Self approved payment",
        },
        {
            "transaction_id": "TXN-R4-01",
            "posting_date": "2026-01-19 13:00:00",
            "amount": 98500.0,
            "vendor_id": "VEND-50",
            "department": "Legal",
            "account_code": "ACC-105",
            "created_by": "grace",
            "approved_by": "heidi",
            "description": "Consulting invoice split",
        },
    ]
    df = pd.DataFrame(data)
    out = io.BytesIO()
    df.to_csv(out, index=False)
    out.seek(0)
    return out


def test_report_export_valid_xlsx_and_expected_sheets(
    client: TestClient,
    auditor_token: str,
    db_session: Session,
):
    """Test generating a valid Excel report with all 5 required sheets and formatting."""
    # 1. Upload dataset
    csv_file = _create_sample_ledger_csv()
    up_res = client.post(
        "/uploads",
        headers={"Authorization": f"Bearer {auditor_token}"},
        files={"file": ("test_report_ledger.csv", csv_file, "text/csv")},
    )
    assert up_res.status_code == 201
    upload_id = up_res.json()["id"]

    # 2. Trigger analysis and execute
    run_res = client.post(
        f"/uploads/{upload_id}/analyze",
        headers={"Authorization": f"Bearer {auditor_token}"},
        json={"settings": {"approval_limit": 100000.0, "flag_threshold": 30}},
    )
    assert run_res.status_code == 202
    run_id = run_res.json()["run_id"]

    # Execute pipeline synchronously
    execute_analysis_pipeline(run_id)

    # 3. Add a review note to one flagged item
    flagged_res = client.get(
        f"/runs/{run_id}/flagged",
        headers={"Authorization": f"Bearer {auditor_token}"},
    )
    assert flagged_res.status_code == 200
    flagged_items = flagged_res.json()["items"]
    assert len(flagged_items) > 0
    first_flagged_id = flagged_items[0]["id"]

    review_res = client.patch(
        f"/flagged/{first_flagged_id}/review",
        headers={"Authorization": f"Bearer {auditor_token}"},
        json={"status": "reviewed", "comment": "Verified against vendor PO #8812."},
    )
    assert review_res.status_code == 200

    # 4. Request Excel report
    report_res = client.get(
        f"/runs/{run_id}/report?format=xlsx",
        headers={"Authorization": f"Bearer {auditor_token}"},
    )
    assert report_res.status_code == 200
    assert (
        report_res.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert (
        f'filename="auditlens_report_run_{run_id}.xlsx"'
        in report_res.headers["content-disposition"]
    )

    # 5. Open workbook with openpyxl and verify all 5 sheets
    wb = openpyxl.load_workbook(io.BytesIO(report_res.content))
    expected_sheets = ["Summary", "Flagged Transactions", "Rule Counts", "Benford", "Review Log"]
    assert wb.sheetnames == expected_sheets

    # Verify Summary sheet content
    ws_sum = wb["Summary"]
    assert "AuditLens" in str(ws_sum["A1"].value)
    # Check that settings and run ID are present
    sum_text = " ".join(str(cell.value) for row in ws_sum.iter_rows() for cell in row if cell.value)
    assert str(run_id) in sum_text
    assert "test_report_ledger.csv" in sum_text
    assert "100,000.00" in sum_text

    # Verify Flagged Transactions sheet content
    ws_flg = wb["Flagged Transactions"]
    assert ws_flg.freeze_panes == "A2"
    assert ws_flg["A1"].value == "Transaction ID"
    assert ws_flg["C1"].value == "Amount (₹)"
    assert ws_flg.max_row > 1  # Headers + flagged items

    # Amount format check
    amount_cell = ws_flg["C2"]
    assert amount_cell.number_format == "#,##0.00"
    assert isinstance(amount_cell.value, (int, float))

    # Verify Rule Counts sheet
    ws_rc = wb["Rule Counts"]
    assert ws_rc["A1"].value == "Rule Code"
    assert ws_rc.max_row >= 7  # Header + R1..R6

    # Verify Benford sheet
    ws_ben = wb["Benford"]
    assert "Benford" in str(ws_ben["A1"].value)

    # Verify Review Log sheet
    ws_rev = wb["Review Log"]
    assert ws_rev["A1"].value == "Logged At (UTC)"
    assert ws_rev["B1"].value == "Transaction ID"
    assert ws_rev.max_row >= 2
    rev_text = " ".join(str(cell.value) for row in ws_rev.iter_rows() for cell in row if cell.value)
    assert "Verified against vendor PO #8812." in rev_text


def test_report_export_zero_flag_case(
    client: TestClient,
    auditor_token: str,
    db_session: Session,
):
    """Test report export when run produces zero flagged exceptions."""
    # 1. Manually create an upload and run with 0 flagged items
    auditor = db_session.query(User).filter_by(role="auditor").first()
    upload = Upload(
        user_id=auditor.id,
        original_filename="clean_ledger.csv",
        stored_path="uploads/test_clean.csv",
        file_type="csv",
        row_count_total=5,
        row_count_accepted=5,
        row_count_rejected=0,
        validation_summary={"total_rows": 5, "valid_rows": 5},
    )
    db_session.add(upload)
    db_session.commit()
    db_session.refresh(upload)

    run = AnalysisRun(
        upload_id=upload.id,
        started_by=auditor.id,
        status="completed",
        settings={"approval_limit": 100000.0, "flag_threshold": 30},
        summary={
            "duration_seconds": 0.42,
            "rules": {"counts_per_rule": {}},
            "scoring": {"total_flagged": 0, "by_level": {}},
        },
        started_at=datetime.now(UTC),
        finished_at=datetime.now(UTC),
    )
    db_session.add(run)
    db_session.commit()
    db_session.refresh(run)

    # 2. Download report
    resp = client.get(
        f"/runs/{run.id}/report?format=xlsx",
        headers={"Authorization": f"Bearer {auditor_token}"},
    )
    assert resp.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == [
        "Summary",
        "Flagged Transactions",
        "Rule Counts",
        "Benford",
        "Review Log",
    ]

    ws_flg = wb["Flagged Transactions"]
    assert "No transactions were flagged" in str(ws_flg["A2"].value)


def test_report_export_formula_injection_case(
    client: TestClient,
    auditor_token: str,
    db_session: Session,
):
    """Test that cell contents starting with =, +, -, @ are neutralized with single quote prefix."""
    auditor = db_session.query(User).filter_by(role="auditor").first()
    upload = Upload(
        user_id=auditor.id,
        original_filename="=HYPERLINK('http://malicious.com','ledger.csv')",
        stored_path="uploads/test_injection.csv",
        file_type="csv",
        row_count_total=1,
        row_count_accepted=1,
        row_count_rejected=0,
        validation_summary={"total_rows": 1, "valid_rows": 1},
    )
    db_session.add(upload)
    db_session.commit()
    db_session.refresh(upload)

    txn = Transaction(
        upload_id=upload.id,
        transaction_id="TXN-INJ-01",
        posting_date=datetime.now(UTC),
        amount=50000.0,
        vendor_id="@DANGEROUS_VENDOR",
        department="+FinanceDept",
        account_code="-ACC99",
        created_by="=cmd|' /C calc'!A0",
        approved_by="+approver",
        description="=SUM(A1:A10)",
    )
    db_session.add(txn)
    db_session.commit()
    db_session.refresh(txn)

    run = AnalysisRun(
        upload_id=upload.id,
        started_by=auditor.id,
        status="completed",
        settings={"approval_limit": 100000.0},
        summary={"duration_seconds": 0.5, "scoring": {"total_flagged": 1}},
        started_at=datetime.now(UTC),
        finished_at=datetime.now(UTC),
    )
    db_session.add(run)
    db_session.commit()
    db_session.refresh(run)

    flagged = FlaggedTransaction(
        run_id=run.id,
        transaction_pk=txn.id,
        risk_score=85,
        risk_level="high",
        rule_points=40,
        ml_points=45,
        triggered_rules=["R1"],
        reasons=["=Dangerous Reason Formula"],
        duplicate_group_id="@DUP_GROUP",
        review_status="open",
    )
    db_session.add(flagged)
    db_session.commit()
    db_session.refresh(flagged)

    create_review(
        db_session,
        flagged_id=flagged.id,
        reviewer_id=auditor.id,
        status="needs_follow_up",
        comment="-Auditor comment starting with minus",
    )

    # Download report
    resp = client.get(
        f"/runs/{run.id}/report?format=xlsx",
        headers={"Authorization": f"Bearer {auditor_token}"},
    )
    assert resp.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))

    # Check Summary sheet filename
    ws_sum = wb["Summary"]
    dataset_name_val = str(ws_sum["B6"].value)
    assert dataset_name_val.startswith("'=")

    # Check Flagged Transactions sheet
    ws_flg = wb["Flagged Transactions"]
    # Vendor ID (col 4) starts with @
    vendor_val = str(ws_flg.cell(row=2, column=4).value)
    assert vendor_val.startswith("'@")

    # Department (col 5) starts with +
    dept_val = str(ws_flg.cell(row=2, column=5).value)
    assert dept_val.startswith("'+")

    # Account code (col 6) starts with -
    acct_val = str(ws_flg.cell(row=2, column=6).value)
    assert acct_val.startswith("'-")

    # Created By (col 7) starts with =
    created_val = str(ws_flg.cell(row=2, column=7).value)
    assert created_val.startswith("'=")

    # Description (col 9) starts with =
    desc_val = str(ws_flg.cell(row=2, column=9).value)
    assert desc_val.startswith("'=")

    # Reasons (col 15) starts with =
    reason_val = str(ws_flg.cell(row=2, column=15).value)
    assert reason_val.startswith("'=")

    # Duplicate Group ID (col 16) starts with @
    dup_val = str(ws_flg.cell(row=2, column=16).value)
    assert dup_val.startswith("'@")

    # Check Review Log sheet
    ws_rev = wb["Review Log"]
    comment_val = str(ws_rev.cell(row=2, column=6).value)
    assert comment_val.startswith("'-")


def test_report_export_ownership_isolation(
    client: TestClient,
    auditor_token: str,
    second_auditor_token: str,
    manager_token: str,
    db_session: Session,
):
    """Test ownership rules: Auditor 2 gets 404 for Auditor 1's report; Manager gets 200."""
    # 1. Auditor 1 uploads and runs analysis
    csv_file = _create_sample_ledger_csv()
    up_res = client.post(
        "/uploads",
        headers={"Authorization": f"Bearer {auditor_token}"},
        files={"file": ("auditor1_ledger.csv", csv_file, "text/csv")},
    )
    upload_id = up_res.json()["id"]

    run_res = client.post(
        f"/uploads/{upload_id}/analyze",
        headers={"Authorization": f"Bearer {auditor_token}"},
    )
    run_id = run_res.json()["run_id"]
    execute_analysis_pipeline(run_id)

    # 2. Auditor 2 attempts to download Auditor 1's report -> 404 Not Found
    res_auditor2 = client.get(
        f"/runs/{run_id}/report?format=xlsx",
        headers={"Authorization": f"Bearer {second_auditor_token}"},
    )
    assert res_auditor2.status_code == 404
    assert res_auditor2.json()["error"]["code"] == "not_found"

    # 3. Manager downloads Auditor 1's report -> 200 OK
    res_mgr = client.get(
        f"/runs/{run_id}/report?format=xlsx",
        headers={"Authorization": f"Bearer {manager_token}"},
    )
    assert res_mgr.status_code == 200


def test_report_export_not_completed_run(
    client: TestClient,
    auditor_token: str,
    db_session: Session,
):
    """Test that requesting report for pending or running run returns 409 Conflict."""
    auditor = db_session.query(User).filter_by(role="auditor").first()
    upload = Upload(
        user_id=auditor.id,
        original_filename="pending_run_ledger.csv",
        stored_path="uploads/pending.csv",
        file_type="csv",
        row_count_total=10,
        row_count_accepted=10,
        row_count_rejected=0,
        validation_summary={},
    )
    db_session.add(upload)
    db_session.commit()
    db_session.refresh(upload)

    run = AnalysisRun(
        upload_id=upload.id,
        started_by=auditor.id,
        status="running",
        settings={},
    )
    db_session.add(run)
    db_session.commit()
    db_session.refresh(run)

    resp = client.get(
        f"/runs/{run.id}/report?format=xlsx",
        headers={"Authorization": f"Bearer {auditor_token}"},
    )
    assert resp.status_code == 409
    assert resp.json()["error"]["code"] == "run_not_completed"


def test_report_export_unsupported_format(
    client: TestClient,
    auditor_token: str,
    db_session: Session,
):
    """Test that requesting unsupported format like pdf or csv returns 422 Unprocessable Entity."""
    auditor = db_session.query(User).filter_by(role="auditor").first()
    upload = Upload(
        user_id=auditor.id,
        original_filename="fmt_test.csv",
        stored_path="uploads/fmt.csv",
        file_type="csv",
        row_count_total=5,
        row_count_accepted=5,
        row_count_rejected=0,
        validation_summary={},
    )
    db_session.add(upload)
    db_session.commit()

    run = AnalysisRun(
        upload_id=upload.id,
        started_by=auditor.id,
        status="completed",
        settings={},
    )
    db_session.add(run)
    db_session.commit()

    resp = client.get(
        f"/runs/{run.id}/report?format=pdf",
        headers={"Authorization": f"Bearer {auditor_token}"},
    )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "unsupported_format"
